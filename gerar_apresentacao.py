#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para gerar uma apresentação PowerPoint em formato carrossel 
do Sistema de Restaurante para LinkedIn
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def criar_slide_titulo(prs, titulo, subtitulo, cor_fundo=None):
    """Cria um slide de título"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Fundo
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(25, 47, 89) if not cor_fundo else cor_fundo
    
    # Título
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(2))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    p.text = titulo
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Subtítulo
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(1.5))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.word_wrap = True
    p = subtitle_frame.paragraphs[0]
    p.text = subtitulo
    p.font.size = Pt(28)
    p.font.color.rgb = RGBColor(100, 181, 246)
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def criar_slide_conteudo(prs, titulo, conteudo, cor_accent=None):
    """Cria um slide com conteúdo"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Fundo
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(240, 242, 245)
    
    # Cor de acento padrão
    if not cor_accent:
        cor_accent = RGBColor(25, 47, 89)
    
    # Barra de título
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = cor_accent
    title_shape.line.color.rgb = cor_accent
    
    # Título
    title_frame = title_shape.text_frame
    title_frame.margin_top = Inches(0.2)
    title_frame.margin_left = Inches(0.4)
    p = title_frame.paragraphs[0]
    p.text = titulo
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    # Conteúdo
    content_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(8.8), Inches(5.5))
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for idx, item in enumerate(conteudo):
        if idx > 0:
            text_frame.add_paragraph()
        p = text_frame.paragraphs[idx]
        p.text = item
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(32, 33, 36)
        p.space_before = Pt(12)
        p.space_after = Pt(12)
        p.level = 0
    
    return slide

def main():
    # Criar apresentação
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Capa
    criar_slide_titulo(
        prs,
        "Sistema de Gestão de Restaurante",
        "Sugestão Inteligente de Compras com .NET e DDD",
        RGBColor(25, 47, 89)
    )
    
    # Slide 2: Visão Geral
    criar_slide_conteudo(
        prs,
        "📋 Visão Geral",
        [
            "✓ Sistema inteligente para gestão de estoque em restaurantes",
            "✓ Análise automática do histórico de vendas dos últimos 7 dias",
            "✓ Calcula ingredientes necessários com base nas receitas",
            "✓ Gera sugestão de compras otimizada",
            "✓ Desenvolvido em C# com .NET e arquitetura DDD"
        ],
        RGBColor(66, 133, 244)
    )
    
    # Slide 3: Tecnologias
    criar_slide_conteudo(
        prs,
        "🛠️ Tecnologias Utilizadas",
        [
            "Linguagem: C# (.NET) - Moderna e type-safe",
            "Banco de Dados: Microsoft SQL Server (Containerizado)",
            "ORM: Entity Framework Core (EF Core)",
            "Arquitetura: Domain-Driven Design (DDD)",
            "Containerização: Docker para ambiente consistente",
            "Padrões: Injeção de Dependência e Migrations"
        ],
        RGBColor(156, 39, 176)
    )
    
    # Slide 4: Arquitetura DDD - Camadas
    criar_slide_conteudo(
        prs,
        "🏛️ Arquitetura em 4 Camadas",
        [
            "1️⃣ Domain: Coração do sistema - Entidades e lógica de negócio",
            "2️⃣ Application: Orquestrador dos casos de uso",
            "3️⃣ Infrastructure: Comunicação com BD e mundo externo",
            "4️⃣ Presentation: Console - Ponto de entrada da aplicação"
        ],
        RGBColor(244, 81, 30)
    )
    
    # Slide 5: Princípios de Engenharia
    criar_slide_conteudo(
        prs,
        "💡 Princípios Aplicados",
        [
            "Separação de Responsabilidades (SoC): Cada camada tem um propósito",
            "Inversão de Dependência (IoC/DI): Componentes acoplados via abstrações",
            "Code-First: Banco de dados modelado diretamente em código",
            "Migrations: Controle de versão do esquema do banco de dados"
        ],
        RGBColor(76, 175, 80)
    )
    
    # Slide 6: Como Executar
    criar_slide_conteudo(
        prs,
        "🚀 Como Executar",
        [
            "1. Subir SQL Server no Docker com um comando",
            "2. Clonar repositório e restaurar dependências",
            "3. Aplicar migrations para criar/atualizar banco de dados",
            "4. Executar: dotnet run",
            "⚡ Pré-requisitos: .NET SDK e Docker instalados"
        ],
        RGBColor(25, 47, 89)
    )
    
    # Slide 7: Funcionalidades Principais
    criar_slide_conteudo(
        prs,
        "⭐ Funcionalidades Principais",
        [
            "📊 Análise inteligente de histórico de vendas",
            "🍽️ Gestão de receitas e ingredientes",
            "📦 Cálculo automático de quantidades necessárias",
            "💰 Otimização de compras baseada em dados reais",
            "🔄 Interface em Console intuitiva e responsiva"
        ],
        RGBColor(244, 127, 32)
    )
    
    # Slide 8: Benefícios
    criar_slide_conteudo(
        prs,
        "🎯 Benefícios",
        [
            "💵 Reduz desperdício e custos de estoque",
            "⏱️ Poupa tempo em análise manual de dados",
            "📈 Decisões baseadas em dados históricos",
            "🔒 Código bem estruturado e fácil de manter",
            "🚀 Escalável e pronto para novos recursos"
        ],
        RGBColor(233, 30, 99)
    )
    
    # Slide 9: Call to Action
    criar_slide_titulo(
        prs,
        "Código Aberto no GitHub",
        "Explore, contribua e use no seu restaurante!",
        RGBColor(76, 175, 80)
    )
    
    # Salvar apresentação
    output_path = "SistemaRestaurante_Apresentacao.pptx"
    prs.save(output_path)
    print(f"✅ Apresentação criada com sucesso: {output_path}")
    print(f"📊 Total de slides: {len(prs.slides)}")

if __name__ == "__main__":
    main()
